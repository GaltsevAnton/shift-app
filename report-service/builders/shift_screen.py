"""
シフト表 — выгрузка экрана シフト管理 (кнопка Excel).

Вид таблицы — как в shift_all.build_range (3 строки на сотрудника: 出勤/退勤/職場).
Отличие только в данных — выводится то, что отфильтровано на экране:
  • сотрудники — ровно те и в том порядке, что на экране (порядок задаёт Java по userIds);
  • колонки №・職種・役職・部署 — по переключателю 表示列 (columns);
  • смены — по фильтру 表示フィルター (workplaces: названия мест + "__none__" + "__off__").
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from models import ShiftScreenRequest, SlotModel
from builders.shift_all import (
    WD_JA, C_OFF_FG, C_OFF_BG, C_ROW_ODD, C_LABEL_BG,
    _thin, _fill, _font, _align, _fmt, _apply_outer_border,
    _date_range, _weekday_from_iso, _day_data_range,
)


def build(req: ShiftScreenRequest) -> bytes:
    dates = _date_range(req.fromDate, req.toDate)
    total = len(dates)
    staff = req.staff

    cols_on = set(req.columns) if req.columns is not None else {"number", "position", "department"}
    wp_on   = set(req.workplaces) if req.workplaces is not None else None   # None = всё видно

    def wp_visible(sl: SlotModel) -> bool:
        if wp_on is None:
            return True
        return (sl.workplace in wp_on) if sl.workplace else ("__none__" in wp_on)

    off_visible = wp_on is None or "__off__" in wp_on

    # Левые колонки (ширины — как в shift_all): [№] [職種・役職] [部署] 氏名 メタ
    left = []
    if "number" in cols_on:
        left.append(("number", "№", 4))
    if "position" in cols_on:
        left.append(("position", "職種・役職", 6))
    if "department" in cols_on:
        left.append(("department", "部署", 8))
    left.append(("name", "氏名", 12))
    left.append(("meta", "日\n曜日", 5))

    n_left    = len(left)
    first_day = n_left + 1          # первая колонка дней
    off_col   = first_day + total   # 公休数
    last_day  = off_col - 1

    wb = Workbook()
    ws = wb.active
    ws.title = f"{req.fromDate}〜{req.toDate}"[:31]

    for i, (_, _, w) in enumerate(left, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(total):
        ws.column_dimensions[get_column_letter(first_day + i)].width = 6
    ws.column_dimensions[get_column_letter(off_col)].width = 5

    ws.row_dimensions[1].height = 6

    ws.merge_cells(f"A2:{get_column_letter(last_day)}2")
    c = ws["A2"]
    c.value     = f"{req.hotelName}シフト表（{req.fromDate}～{req.toDate}）"
    c.font      = _font(bold=True, size=12)
    c.alignment = _align(h="center")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 5

    for col_idx, (_, label, _) in enumerate(left, start=1):
        ws.merge_cells(f"{get_column_letter(col_idx)}4:{get_column_letter(col_idx)}5")
        c = ws.cell(row=4, column=col_idx)
        c.value     = label
        c.font      = _font(size=8)
        c.alignment = _align()
        c.border    = _thin()
        _apply_outer_border(ws, 4, 5, col_idx, col_idx)

    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 16

    for i, ds in enumerate(dates):
        col = first_day + i
        wd  = _weekday_from_iso(ds)
        dd  = int(ds[8:10])

        c = ws.cell(row=4, column=col)
        c.value     = dd
        c.font      = _font(size=8)
        c.alignment = _align()
        c.border    = _thin()

        c = ws.cell(row=5, column=col)
        c.value     = WD_JA[wd]
        c.font      = _font(size=8)
        c.alignment = _align()
        c.border    = _thin()

    ws.merge_cells(f"{get_column_letter(off_col)}4:{get_column_letter(off_col)}5")
    c = ws.cell(row=4, column=off_col)
    c.value     = "公休\n数"
    c.font      = _font(size=7)
    c.alignment = _align()
    c.border    = _thin()
    _apply_outer_border(ws, 4, 5, off_col, off_col)

    base_row = 6
    for staff_idx, s in enumerate(staff):
        r      = base_row + staff_idx * 3
        row_bg = C_ROW_ODD

        for col_idx, (key, _, _) in enumerate(left, start=1):
            if key == "meta":
                for sub, label in enumerate(["出勤", "退勤", "職場"]):
                    c = ws.cell(row=r + sub, column=col_idx)
                    c.value     = label
                    c.font      = _font(size=7, color="333333")
                    c.fill      = _fill(C_LABEL_BG)
                    c.alignment = _align()
                    c.border    = _thin()
                    ws.row_dimensions[r + sub].height = 14
                continue

            col_l = get_column_letter(col_idx)
            ws.merge_cells(f"{col_l}{r}:{col_l}{r+2}")
            c = ws[f"{col_l}{r}"]
            if key == "number":
                c.value = s.sortOrder if s.sortOrder is not None else ""
                c.font  = _font(size=8, color="555555")
            elif key == "position":
                c.value = s.position or ""
                c.font  = _font(size=8, color="555555")
            elif key == "department":
                c.value = "、".join(s.departments)
                c.font  = _font(size=8, color="555555")
            else:  # name
                c.value = s.userName
                c.font  = _font(bold=True, size=9)
            c.fill      = _fill(row_bg)
            c.alignment = _align(h="center")
            c.border    = _thin()
            _apply_outer_border(ws, r, r+2, col_idx, col_idx)

        for i, ds in enumerate(dates):
            col = first_day + i
            day_bg = row_bg

            d = _day_data_range(s, ds)

            if d.off or not d.slots:
                ws.merge_cells(f"{get_column_letter(col)}{r}:{get_column_letter(col)}{r+2}")
                c = ws.cell(row=r, column=col)
                if off_visible:
                    c.value = "休"
                    c.font  = _font(color=C_OFF_FG, bold=True, size=10)
                    c.fill  = _fill(C_OFF_BG)
                c.alignment = _align()
                c.border    = _thin()
                _apply_outer_border(ws, r, r+2, col, col)
            else:
                slots = [sl for sl in d.slots if wp_visible(sl)]   # 表示フィルター
                starts     = "\n".join(_fmt(sl.startTime) for sl in slots if sl.startTime)
                ends       = "\n".join("L" if sl.last else _fmt(sl.endTime) for sl in slots)
                workplaces = "\n".join(sl.workplace or "" for sl in slots)

                for sub, val in enumerate([starts, ends, workplaces]):
                    c = ws.cell(row=r + sub, column=col)
                    c.value     = val
                    c.fill      = _fill(day_bg)
                    c.alignment = _align(h="center", v="center")
                    c.border    = _thin()
                    c.font      = _font(size=9 if sub in (0, 1) else 8,
                                        color="555555" if sub == 2 else "000000")

        # 公休数 — по всем дням (как на экране)
        off_count = sum(1 for ds in dates if not _day_data_range(s, ds).slots or _day_data_range(s, ds).off)
        ws.merge_cells(f"{get_column_letter(off_col)}{r}:{get_column_letter(off_col)}{r+2}")
        c = ws.cell(row=r, column=off_col)
        c.value     = off_count
        c.font      = _font(size=9)
        c.alignment = _align()
        c.border    = _thin()
        _apply_outer_border(ws, r, r+2, off_col, off_col)

        thick = Side(style="medium", color="000000")
        no    = Side(style=None)
        is_first = staff_idx == 0

        for sub_row in range(r, r + 3):
            for col_idx in range(1, off_col):
                c = ws.cell(row=sub_row, column=col_idx)
                top    = thick if (sub_row == r and is_first) else no
                bottom = thick if sub_row == r + 2 else no
                lft    = thick if col_idx == 1 else no
                right  = thick if col_idx == last_day else no
                if any([top != no, bottom != no, lft != no, right != no]):
                    existing = c.border
                    c.border = Border(
                        top    = top    if top    != no else existing.top,
                        bottom = bottom if bottom != no else existing.bottom,
                        left   = lft    if lft    != no else existing.left,
                        right  = right  if right  != no else existing.right,
                    )

    last_data_row = base_row + len(staff) * 3 - 1

    thick = Side(style="medium", color="000000")
    no    = Side(style=None)

    for row in range(4, last_data_row + 1):
        for col in range(1, off_col + 1):
            c = ws.cell(row=row, column=col)
            top    = thick if row == 4             else no
            bottom = thick if row == last_data_row else no
            lft    = thick if col == 1             else no
            right  = thick if col == off_col       else no
            if any([top, bottom, lft, right]):
                existing = c.border
                c.border = Border(
                    top    = top    if top    != no else existing.top,
                    bottom = bottom if bottom != no else existing.bottom,
                    left   = lft    if lft    != no else existing.left,
                    right  = right  if right  != no else existing.right,
                )

    ws.freeze_panes = ws.cell(row=6, column=first_day)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize   = 9
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.75, bottom=0.75, header=0.3, footer=0.3
    )
    ws.page_setup.copies = 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()